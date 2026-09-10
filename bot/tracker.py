#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Booking Tracker — Maintains Excel file + formatted Telegram summaries
"""

import json
import os
from datetime import datetime

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
BOOKINGS_FILE = os.path.join(SCRIPT_DIR, "bookings.json")

def load_bookings():
    if os.path.exists(BOOKINGS_FILE):
        with open(BOOKINGS_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    return []

def get_all_bookings():
    return load_bookings()

def get_pending_bookings():
    return [b for b in load_bookings() if b.get("status") == "pending"]

def get_confirmed_bookings():
    return [b for b in load_bookings() if b.get("status") == "confirmed"]

def format_booking_message(booking, lang="ar"):
    """Format a single booking for Telegram"""
    status_emoji = {"pending": "⏳", "confirmed": "✅", "rejected": "❌"}.get(booking.get("status"), "❓")
    status_text = {"pending": "قيد الانتظار", "confirmed": "مؤكد", "rejected": "مرفوض"}.get(booking.get("status"), "غير معروف")
    
    msg = (
        f"{status_emoji} **{booking.get('id', 'N/A')}** — {status_text}\n"
        f"  👤 {booking.get('name', 'N/A')}\n"
        f"  📱 {booking.get('phone', 'N/A')}\n"
        f"  💼 {booking.get('type_name', 'N/A')}\n"
        f"  📅 {booking.get('date', 'N/A')} — ⏰ {booking.get('time', 'N/A')}\n"
        f"  📝 {booking.get('details', 'لا توجد تفاصيل')}\n"
    )
    return msg

def format_summary_message(lang="ar"):
    """Format all bookings as a summary"""
    bookings = load_bookings()
    if not bookings:
        return "📋 لا توجد حجوزات بعد." if lang == "ar" else "📋 No bookings yet."
    
    pending = [b for b in bookings if b.get("status") == "pending"]
    confirmed = [b for b in bookings if b.get("status") == "confirmed"]
    rejected = [b for b in bookings if b.get("status") == "rejected"]
    
    msg = "📊 **ملخص الحجوزات**\n━━━━━━━━━━━━━━━━━━━━\n"
    msg += f"⏳ قيد الانتظار: {len(pending)}\n"
    msg += f"✅ مؤكد: {len(confirmed)}\n"
    msg += f"❌ مرفوض: {len(rejected)}\n"
    msg += f"📊 الإجمالي: {len(bookings)}\n"
    msg += "━━━━━━━━━━━━━━━━━━━━\n\n"
    
    if pending:
        msg += "⏳ **قيد الانتظار:**\n"
        for b in pending:
            msg += format_booking_message(b) + "\n"
    
    if confirmed:
        msg += "\n✅ **مؤكد:**\n"
        for b in confirmed:
            msg += format_booking_message(b) + "\n"
    
    if rejected:
        msg += "\n❌ **مرفوض:**\n"
        for b in rejected:
            msg += format_booking_message(b) + "\n"
    
    return msg

def create_excel_report():
    """Create an Excel file with all bookings (timestamped name → never collides with open files)"""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        # If openpyxl not installed, use simple CSV
        return create_csv_report()
    
    bookings = load_bookings()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bookings"
    
    # Styles
    header_font = Font(bold=True, color="FFFFFF", size=12)
    header_fill = PatternFill(start_color="D4A843", end_color="D4A843", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin")
    )
    
    # Headers
    headers = ["رقم الحجز", "الاسم", "الهاتف", "النوع", "التاريخ", "الوقت", "التفاصيل", "الحالة", "تاريخ الإنشاء"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = border
    
    # Data
    status_map = {"pending": "قيد الانتظار", "confirmed": "مؤكد", "rejected": "مرفوض"}
    for row, booking in enumerate(bookings, 2):
        data = [
            booking.get("id", ""),
            booking.get("name", ""),
            booking.get("phone", ""),
            booking.get("type_name", ""),
            booking.get("date", ""),
            booking.get("time", ""),
            booking.get("details", ""),
            status_map.get(booking.get("status", ""), booking.get("status", "")),
            booking.get("timestamp", "")[:10] if booking.get("timestamp") else ""
        ]
        for col, value in enumerate(data, 1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.border = border
            cell.alignment = Alignment(vertical="center")
    
    # Auto-width
    for col in ws.columns:
        max_length = 0
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        ws.column_dimensions[col[0].column_letter].width = min(max_length + 2, 30)
    
    # Timestamped filename → never blocked by an open copy in Excel
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filepath = os.path.join(SCRIPT_DIR, f"bookings_report_{stamp}.xlsx")
    wb.save(filepath)

    # Also save a "latest" copy with a FIXED name → always openable from the machine
    latest_path = os.path.join(SCRIPT_DIR, "bookings_report_latest.xlsx")
    try:
        wb.save(latest_path)
    except Exception:
        pass  # if it's open in Excel, the timestamped copy still works

    return filepath

def create_csv_report():
    """Fallback: Create CSV file"""
    import csv
    bookings = load_bookings()
    filepath = os.path.join(SCRIPT_DIR, "bookings_report.csv")
    
    with open(filepath, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["رقم الحجز", "الاسم", "الهاتف", "النوع", "التاريخ", "الوقت", "التفاصيل", "الحالة"])
        status_map = {"pending": "قيد الانتظار", "confirmed": "مؤكد", "rejected": "مرفوض"}
        for b in bookings:
            writer.writerow([
                b.get("id", ""),
                b.get("name", ""),
                b.get("phone", ""),
                b.get("type_name", ""),
                b.get("date", ""),
                b.get("time", ""),
                b.get("details", ""),
                status_map.get(b.get("status", ""), b.get("status", ""))
            ])
    return filepath
